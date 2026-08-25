import logging
import re
from copy import deepcopy
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE

from hopptx.schemas.metrics import Metrics

log = logging.getLogger(__name__)


def _format_decimal(value: Decimal) -> str:
    """Format decimals without scientific notation and trim trailing zeros."""
    value_str = format(value, "f")
    if "." not in value_str:
        return value_str
    return value_str.rstrip("0").rstrip(".")


def _format_currency(value: Decimal) -> str:
    """Format currency with thousands separators and two decimal places."""
    return format(value, ",.2f")


def _format_date_range(start: str, end: str) -> str:
    """Format date range as human-readable period (e.g., 'January – March 2026')."""
    start_dt = datetime.strptime(start, "%Y-%m-%d")
    end_dt = datetime.strptime(end, "%Y-%m-%d")

    start_month = start_dt.strftime("%B")
    end_month = end_dt.strftime("%B")
    start_year = start_dt.strftime("%Y")
    end_year = end_dt.strftime("%Y")

    if start_month == end_month and start_year == end_year:
        return f"{start_month} {start_year}"
    if start_year == end_year:
        return f"{start_month} – {end_month} {start_year}"
    else:
        return f"{start_month} {start_year} – {end_month} {end_year}"


def _replace_placeholders_in_text(text: str, replacements: dict[str, str]) -> str:
    """Replace all {{placeholder}} patterns in text."""
    for key, value in replacements.items():
        placeholder = f"{{{{{key}}}}}"
        text = text.replace(placeholder, value)
    return text


def _replace_placeholders_in_run(run, replacements: dict[str, str]) -> str:
    """Replace placeholders in a single run, preserving its formatting."""
    return _replace_placeholders_in_text(run.text, replacements)


def _replace_placeholders_in_paragraph(paragraph, replacements: dict[str, str]) -> None:
    """Replace placeholders in a paragraph while preserving run formatting.

    Strategy:
    1. First, try replacing placeholders within individual runs (fast path)
    2. If placeholders span multiple runs, concatenate all runs, replace,
       then recreate with original first run's formatting
    """
    # Collect all runs with their formatting
    runs = list(paragraph.runs)
    if not runs:
        return

    # Get full paragraph text
    full_text = paragraph.text
    if not full_text:
        return

    # Check if any placeholder exists
    has_placeholder = any(f"{{{{{key}}}}}" in full_text for key in replacements.keys())
    if not has_placeholder:
        return

    # First, try replacing in each run individually
    for run in runs:
        new_run_text = _replace_placeholders_in_run(run, replacements)
        if new_run_text != run.text:
            run.text = new_run_text

    # Check if placeholders remain (they were split across runs)
    remaining_text = paragraph.text
    has_remaining = any(
        f"{{{{{key}}}}}" in remaining_text for key in replacements.keys()
    )

    if has_remaining:
        # Merge all runs into one to handle split placeholders
        new_text = _replace_placeholders_in_text(remaining_text, replacements)
        if new_text != remaining_text:
            first_run = runs[0]
            # Clear and recreate with first run's formatting
            paragraph.clear()
            new_run = paragraph.add_run()
            new_run.text = new_text
            _copy_run_font(first_run, new_run)


def _collect_all_text(prs: Presentation) -> str:
    """Collect all text from the presentation for placeholder detection."""
    parts: list[str] = []
    for slide in prs.slides:
        for shape in _iter_shapes(slide.shapes):
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        parts.append(cell.text_frame.text)
            elif shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


PROGRAMS_COLUMN_KEY_RE = re.compile(
    r"^programs_(?P<category>[a-z0-9_]+)_col_(?P<col>\d+)$"
)
PROGRAMS_TABLE_ANCHOR_RE = re.compile(r"\{\{programs_table_p(?P<page>\d+)\}\}")
PROGRAMS_TABLE_CATEGORY_ANCHOR_RE = re.compile(
    r"\{\{programs_table_category_p(?P<page>\d+)\}\}"
)
PROGRAMS_TABLE_TOPIC_ANCHOR_RE = re.compile(
    r"\{\{programs_table_topic_p(?P<page>\d+)\}\}"
)
TOP_BOTTOM_ROW_PLACEHOLDER_RE = re.compile(
    r"\{\{top_(?P<index>\d+)_(?:name|count)\}\}"
)
TOP_PROGRAMS_RANK_PLACEHOLDER = "{{n_rank}}"
TOP_PROGRAMS_NAME_PLACEHOLDER = "{{top_n_programs_name}}"
TOP_PROGRAMS_COUNT_PLACEHOLDER = "{{top_n_programs_count}}"
PROGRAMS_TABLE_CATEGORY_ORDER = (
    "Technical Skills",
    "Business Skills",
    "Leading Others",
    "Communication Skills",
    "Influencing Skills",
    "Personal / Professional Development",
)
_PROGRAMS_TABLE_CATEGORY_RANK = {
    name.casefold(): index for index, name in enumerate(PROGRAMS_TABLE_CATEGORY_ORDER)
}


def _iter_shapes(shapes):
    """Yield shapes recursively, including children of grouped shapes."""
    for shape in shapes:
        yield shape
        child_shapes = getattr(shape, "shapes", None)
        if child_shapes is not None:
            yield from _iter_shapes(child_shapes)


@dataclass(frozen=True)
class _ProgramsTableAnchor:
    page: int
    slide: object
    table: object
    anchor_row: int
    category_col: int
    topic_start_col: int


@dataclass(frozen=True)
class _TopProgramsTableAnchor:
    table: object
    anchor_row: int
    rank_col: int
    name_col: int
    count_col: int


def _slugify_identifier(text: str) -> str:
    """Normalize text into snake_case identifier form for placeholder matching."""
    slug = re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")
    return re.sub(r"_+", "_", slug)


def _extract_template_placeholders(prs: Presentation) -> set[str]:
    all_text = _collect_all_text(prs)
    return set(re.findall(r"\{\{(\w+)\}\}", all_text))


def _split_evenly(items: list[str], num_columns: int) -> list[list[str]]:
    """Split items into balanced contiguous chunks across num_columns."""
    if num_columns <= 0:
        raise ValueError(f"num_columns must be positive, got {num_columns}")
    base = len(items) // num_columns
    remainder = len(items) % num_columns
    chunks: list[list[str]] = []
    start = 0
    for i in range(num_columns):
        size = base + (1 if i < remainder else 0)
        end = start + size
        chunks.append(items[start:end])
        start = end
    return chunks


def _sort_programs_table_categories(
    categories: list[tuple[str, list[str]]],
) -> list[tuple[str, list[str]]]:
    """Order categories for the programs table using a fixed business-defined sequence."""
    ordered: list[tuple[int, str, list[str]]] = []
    overflow: list[tuple[str, list[str]]] = []

    for category_name, programs in categories:
        rank = _PROGRAMS_TABLE_CATEGORY_RANK.get(category_name.strip().casefold())
        if rank is None:
            overflow.append((category_name, programs))
            continue
        ordered.append((rank, category_name, programs))

    ordered.sort(key=lambda item: item[0])
    return [(name, programs) for _, name, programs in ordered] + overflow


def _truncate_text(text: str, max_chars: int | None, suffix: str | None) -> str:
    """Trim text to max_chars and append suffix when trimmed.

    If suffix is None, truncation is disabled and the original text is returned.
    """
    if suffix is None:
        return text
    if max_chars is None or max_chars <= 0:
        raise ValueError(
            f"max_chars must be a positive int when truncation is enabled, got {max_chars}"
        )
    if len(text) <= max_chars:
        return text
    keep = max_chars - len(suffix)
    if keep <= 0:
        return suffix[:max_chars]
    return text[:keep].rstrip() + suffix


def _build_program_column_replacements(
    programs_by_category: dict[str, list[str]],
    template_placeholders: set[str],
    topic_max_chars: int | None,
    topic_truncation_suffix: str | None,
) -> dict[str, str]:
    """Build replacements for programs_<category>_col_<n> placeholders."""
    placeholders_by_category: dict[str, list[tuple[int, str]]] = {}
    for key in template_placeholders:
        match = PROGRAMS_COLUMN_KEY_RE.match(key)
        if match is None:
            continue
        category = match.group("category")
        col_num = int(match.group("col"))
        placeholders_by_category.setdefault(category, []).append((col_num, key))

    if not placeholders_by_category:
        return {}

    categories_by_slug: dict[str, list[str]] = {}
    for category_name, items in programs_by_category.items():
        slug = _slugify_identifier(category_name)
        if slug in categories_by_slug:
            raise ValueError(
                f"Category slug collision for '{slug}'. "
                f"At least two categories normalize to the same placeholder key."
            )
        categories_by_slug[slug] = items
    replacements: dict[str, str] = {}

    for category_slug, keys in placeholders_by_category.items():
        ordered_keys = [key for _, key in sorted(keys, key=lambda x: x[0])]
        items = categories_by_slug.get(category_slug, [])
        items = [
            _truncate_text(item, topic_max_chars, topic_truncation_suffix)
            for item in items
        ]
        chunks = _split_evenly(items, len(ordered_keys))
        for key, chunk in zip(ordered_keys, chunks):
            replacements[key] = "\n".join(chunk)

    missing_in_template = sorted(
        category
        for category in programs_by_category
        if _slugify_identifier(category) not in placeholders_by_category
    )
    if missing_in_template:
        log.warning(
            f"{len(missing_in_template)} category/categor(ies) from data have no matching "
            f"programs_<category>_col_<n> placeholder(s): {missing_in_template}"
        )

    return replacements


def _find_programs_table_anchors(prs: Presentation) -> list[_ProgramsTableAnchor]:
    anchors_by_page: dict[int, _ProgramsTableAnchor] = {}
    category_hits: dict[int, tuple[object, object, int, int]] = {}
    topic_hits: dict[int, tuple[object, object, int, int]] = {}

    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            table = shape.table
            for row_idx, row in enumerate(table.rows):
                for col_idx, cell in enumerate(row.cells):
                    text = cell.text_frame.text
                    category_match = PROGRAMS_TABLE_CATEGORY_ANCHOR_RE.search(text)
                    if category_match is not None:
                        page = int(category_match.group("page"))
                        if page in category_hits:
                            raise ValueError(
                                f"Found multiple '{{{{programs_table_category_p{page}}}}}' anchors in template."
                            )
                        category_hits[page] = (slide, table, row_idx, col_idx)

                    topic_match = PROGRAMS_TABLE_TOPIC_ANCHOR_RE.search(text)
                    if topic_match is not None:
                        page = int(topic_match.group("page"))
                        if page in topic_hits:
                            raise ValueError(
                                f"Found multiple '{{{{programs_table_topic_p{page}}}}}' anchors in template."
                            )
                        topic_hits[page] = (slide, table, row_idx, col_idx)

                    match = PROGRAMS_TABLE_ANCHOR_RE.search(text)
                    if match is None:
                        continue
                    page = int(match.group("page"))
                    if page in anchors_by_page:
                        raise ValueError(
                            f"Found multiple '{{{{programs_table_p{page}}}}}' anchors in template. "
                            f"Keep exactly one table anchor per page number."
                        )
                    anchors_by_page[page] = _ProgramsTableAnchor(
                        page=page,
                        slide=slide,
                        table=table,
                        anchor_row=row_idx,
                        category_col=0,
                        topic_start_col=1,
                    )

    if category_hits or topic_hits:
        pages = sorted(set(category_hits) | set(topic_hits))
        for page in pages:
            if page not in category_hits or page not in topic_hits:
                raise ValueError(
                    f"Template page {page} must include both "
                    f"'{{{{programs_table_category_p{page}}}}}' and '{{{{programs_table_topic_p{page}}}}}'."
                )
            slide_c, table_c, row_c, col_c = category_hits[page]
            slide_t, table_t, row_t, col_t = topic_hits[page]
            if table_c is not table_t or row_c != row_t:
                raise ValueError(
                    f"Template page {page} anchors must be in the same table row."
                )
            if col_t <= col_c:
                raise ValueError(
                    f"Template page {page} topic anchor must be to the right of category anchor."
                )
            anchors_by_page[page] = _ProgramsTableAnchor(
                page=page,
                slide=slide_c,
                table=table_c,
                anchor_row=row_c,
                category_col=col_c,
                topic_start_col=col_t,
            )

    return [anchors_by_page[page] for page in sorted(anchors_by_page)]


def _copy_run_font(src_run, dst_run) -> None:
    if src_run is None:
        return
    if src_run.font.size:
        dst_run.font.size = src_run.font.size
    if src_run.font.name:
        dst_run.font.name = src_run.font.name
    src_color = src_run.font.color
    dst_color = dst_run.font.color
    if src_color.type == MSO_COLOR_TYPE.RGB and src_color.rgb is not None:
        dst_color.rgb = src_color.rgb
    elif (
        src_color.type == MSO_COLOR_TYPE.SCHEME
        and src_color.theme_color is not None
    ):
        dst_color.theme_color = src_color.theme_color
    if src_color.type is not None and src_color.brightness != 0:
        dst_color.brightness = src_color.brightness
    if src_run.font.bold is not None:
        dst_run.font.bold = src_run.font.bold
    if src_run.font.italic is not None:
        dst_run.font.italic = src_run.font.italic
    if src_run.font.underline is not None:
        dst_run.font.underline = src_run.font.underline


def _set_table_cell_lines(
    cell,
    lines: list[str],
    bullet: bool,
    font_source_run=None,
    paragraph_source=None,
) -> None:
    if not lines:
        cell.text = ""
        return
    cell.text = ""
    text_frame = cell.text_frame
    for idx, line in enumerate(lines):
        para = text_frame.paragraphs[0] if idx == 0 else text_frame.add_paragraph()
        if paragraph_source is not None:
            para.alignment = paragraph_source.alignment
        run = para.add_run()
        run.text = f"• {line}" if bullet else line
        _copy_run_font(font_source_run, run)


def _strip_ext_lists(tr_element) -> None:
    """Remove extension lists from a copied table row to prevent frozen cells.

    Deep-copied rows carry ``<a:extLst>`` elements that contain unique creation
    IDs (``a16:creationId`` etc.).  When multiple cells share the same IDs,
    PowerPoint treats the duplicates as non-editable.  Stripping the extension
    lists from cells and their text bodies makes the rows fully editable.
    """
    nsmap = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    for ext_lst in tr_element.findall(".//a:extLst", nsmap):
        parent = ext_lst.getparent()
        if parent is not None:
            parent.remove(ext_lst)


def _ensure_table_rows(table, anchor_row: int, required_count: int) -> None:
    """Ensure table has at least `required_count` rows available from `anchor_row` down."""
    if required_count < 0:
        raise ValueError(f"required_count must be non-negative, got {required_count}")
    existing_count = len(table.rows) - anchor_row
    if existing_count >= required_count:
        return

    if not table.rows:
        raise ValueError("Programs table has no rows")
    ref_idx = anchor_row if anchor_row < len(table.rows) else len(table.rows) - 1
    template_tr = table._tbl.tr_lst[ref_idx]
    for _ in range(required_count - existing_count):
        new_tr = deepcopy(template_tr)
        _strip_ext_lists(new_tr)
        table._tbl.append(new_tr)


def _remove_slide(prs: Presentation, slide) -> None:
    """Remove a slide from the presentation."""
    slides = prs.slides
    slide_list = list(slides)
    if slide not in slide_list:
        return
    idx = slide_list.index(slide)
    r_id = slides._sldIdLst[idx].rId
    prs.part.drop_rel(r_id)
    del slides._sldIdLst[idx]


def _prune_top_table_rows(prs: Presentation, top_count: int) -> None:
    """Remove table rows that only contain out-of-range top placeholders."""
    if top_count < 0:
        raise ValueError("top_count must be non-negative")

    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_table:
                continue

            table = shape.table
            rows_to_remove: list[int] = []

            for row_idx, row in enumerate(table.rows):
                row_text = " ".join(cell.text_frame.text for cell in row.cells)
                matches = list(TOP_BOTTOM_ROW_PLACEHOLDER_RE.finditer(row_text))
                if not matches:
                    continue

                keep_row = False
                for match in matches:
                    index = int(match.group("index"))
                    if index <= top_count:
                        keep_row = True
                        break

                if not keep_row:
                    rows_to_remove.append(row_idx)

            for row_idx in reversed(rows_to_remove):
                table._tbl.remove(table._tbl.tr_lst[row_idx])


def _find_top_programs_table_anchor(prs: Presentation) -> _TopProgramsTableAnchor | None:
    anchor: _TopProgramsTableAnchor | None = None
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            table = shape.table
            for row_idx, row in enumerate(table.rows):
                rank_col = None
                name_col = None
                count_col = None
                for col_idx, cell in enumerate(row.cells):
                    text = cell.text_frame.text
                    if TOP_PROGRAMS_RANK_PLACEHOLDER in text:
                        if rank_col is not None:
                            raise ValueError(
                                "Found multiple '{{n_rank}}' placeholders in a single row."
                            )
                        rank_col = col_idx
                    if TOP_PROGRAMS_NAME_PLACEHOLDER in text:
                        if name_col is not None:
                            raise ValueError(
                                "Found multiple '{{top_n_programs_name}}' placeholders in a single row."
                            )
                        name_col = col_idx
                    if TOP_PROGRAMS_COUNT_PLACEHOLDER in text:
                        if count_col is not None:
                            raise ValueError(
                                "Found multiple '{{top_n_programs_count}}' placeholders in a single row."
                            )
                        count_col = col_idx

                if (
                    rank_col is not None
                    or name_col is not None
                    or count_col is not None
                ):
                    if None in (rank_col, name_col, count_col):
                        raise ValueError(
                            "Top programs table row must include all placeholders: "
                            "'{{n_rank}}', '{{top_n_programs_name}}', '{{top_n_programs_count}}'."
                        )
                    if anchor is not None:
                        raise ValueError(
                            "Found multiple top programs table anchors. "
                            "Keep exactly one row with '{{n_rank}}', '{{top_n_programs_name}}', "
                            "and '{{top_n_programs_count}}'."
                        )
                    anchor = _TopProgramsTableAnchor(
                        table=table,
                        anchor_row=row_idx,
                        rank_col=rank_col,
                        name_col=name_col,
                        count_col=count_col,
                    )
    return anchor


def _rank_top_topics(top_topics: list[tuple[str, int]]) -> list[tuple[int, str, int]]:
    ranked: list[tuple[int, str, int]] = []
    current_rank = 0
    last_count: int | None = None
    for topic, count in top_topics:
        if count != last_count:
            current_rank += 1
            last_count = count
        ranked.append((current_rank, topic, count))
    return ranked


def _populate_top_programs_table(prs: Presentation, top_topics: list[tuple[str, int]]) -> bool:
    """Fill a top programs table anchored by n-rank placeholders."""
    anchor = _find_top_programs_table_anchor(prs)
    if anchor is None:
        return False

    ranked_topics = _rank_top_topics(top_topics)
    if not ranked_topics:
        raise ValueError("Cannot populate top programs table: no topics available.")

    _ensure_table_rows(anchor.table, anchor.anchor_row, len(ranked_topics))
    rank_font_source = None
    name_font_source = None
    count_font_source = None
    rank_paragraph_source = None
    name_paragraph_source = None
    count_paragraph_source = None
    if anchor.anchor_row < len(anchor.table.rows):
        rank_paragraph_source = anchor.table.cell(
            anchor.anchor_row, anchor.rank_col
        ).text_frame.paragraphs[0]
        rank_runs = list(rank_paragraph_source.runs)
        if rank_runs:
            rank_font_source = rank_runs[0]
        name_paragraph_source = anchor.table.cell(
            anchor.anchor_row, anchor.name_col
        ).text_frame.paragraphs[0]
        name_runs = list(name_paragraph_source.runs)
        if name_runs:
            name_font_source = name_runs[0]
        count_paragraph_source = anchor.table.cell(
            anchor.anchor_row, anchor.count_col
        ).text_frame.paragraphs[0]
        count_runs = list(count_paragraph_source.runs)
        if count_runs:
            count_font_source = count_runs[0]

    for row_idx in range(anchor.anchor_row, len(anchor.table.rows)):
        for col_idx in range(len(anchor.table.columns)):
            anchor.table.cell(row_idx, col_idx).text = ""

    for idx, (rank, topic, count) in enumerate(ranked_topics):
        row_idx = anchor.anchor_row + idx
        _set_table_cell_lines(
            anchor.table.cell(row_idx, anchor.rank_col),
            [str(rank)],
            bullet=False,
            font_source_run=rank_font_source,
            paragraph_source=rank_paragraph_source,
        )
        _set_table_cell_lines(
            anchor.table.cell(row_idx, anchor.name_col),
            [topic],
            bullet=False,
            font_source_run=name_font_source,
            paragraph_source=name_paragraph_source,
        )
        _set_table_cell_lines(
            anchor.table.cell(row_idx, anchor.count_col),
            [str(count)],
            bullet=False,
            font_source_run=count_font_source,
            paragraph_source=count_paragraph_source,
        )

    required_rows = anchor.anchor_row + len(ranked_topics)
    while len(anchor.table.rows) > required_rows:
        anchor.table._tbl.remove(anchor.table._tbl.tr_lst[-1])

    return True


def _paginate_categories_by_program_limit(
    categories: list[tuple[str, list[str]]],
    anchors: list[_ProgramsTableAnchor],
    max_programs_per_table_page: int,
    single_page_soft_overflow: int,
) -> list[list[tuple[str, list[str]]]]:
    """Assign whole categories to pages using a greedy fill strategy.

    Packs categories onto the first page up to the hard per-page limit, then
    overflows onto the next page. Earlier pages are filled fuller; later
    pages take whatever is left over. A soft-overflow window keeps near-limit
    totals on a single page.
    """
    per_page: list[list[tuple[str, list[str]]]] = [[] for _ in anchors]
    if not categories:
        log.info("No categories to paginate.")
        return per_page

    total_programs = sum(max(1, len(programs)) for _, programs in categories)
    log.info(
        f"Paginating {total_programs} program(s) across {len(categories)} "
        f"categor(ies) into up to {len(anchors)} page(s) "
        f"(limit {max_programs_per_table_page}/page, "
        f"soft overflow +{single_page_soft_overflow})."
    )

    soft_limit = max_programs_per_table_page + single_page_soft_overflow
    if total_programs <= soft_limit:
        per_page[0] = categories
        log.info(
            f"All {total_programs} program(s) fit on page 1 "
            f"(within soft limit of {soft_limit}); skipping split."
        )
        _log_pagination_result(per_page)
        return per_page

    # Greedy fill: fit as many whole categories as possible on the current
    # page up to the hard limit, then move to the next page.
    page_idx = 0
    page_total = 0
    for name, programs in categories:
        count = max(1, len(programs))
        if per_page[page_idx] and (page_total + count) > max_programs_per_table_page:
            page_idx += 1
            page_total = 0
            if page_idx >= len(anchors):
                raise ValueError(
                    f"Not enough programs table pages at "
                    f"{max_programs_per_table_page} programs/page. "
                    f"Add '{{{{programs_table_p{len(anchors) + 1}}}}}' to a continuation slide."
                )
        per_page[page_idx].append((name, programs))
        page_total += count

    log.info(f"Split across {page_idx + 1} page(s) using greedy fill.")
    _log_pagination_result(per_page)
    return per_page


def _log_pagination_result(
    per_page: list[list[tuple[str, list[str]]]],
) -> None:
    """Log per-page breakdown of categories and program counts."""
    for page_idx, page_categories in enumerate(per_page, start=1):
        if not page_categories:
            log.info(f"  Page {page_idx}: empty (slide will be removed)")
            continue
        page_total = sum(max(1, len(programs)) for _, programs in page_categories)
        log.info(
            f"  Page {page_idx}: {len(page_categories)} categor(ies), "
            f"{page_total} program(s)"
        )
        for name, programs in page_categories:
            log.info(f"    - {name}: {len(programs)} program(s)")


def _populate_programs_tables(
    prs: Presentation,
    programs_by_category: dict[str, list[str]],
    max_programs_per_table_page: int,
    single_page_soft_overflow: int,
    topic_max_chars: int | None,
    topic_truncation_suffix: str | None,
) -> bool:
    """Fill anchored programs tables page-by-page without splitting categories."""
    anchors = _find_programs_table_anchors(prs)
    if not anchors:
        return False

    program_cols_by_page: dict[int, int] = {}

    for anchor in anchors:
        num_cols = len(anchor.table.columns)
        if num_cols < anchor.topic_start_col + 1:
            raise ValueError(
                f"programs_table_p{anchor.page} must have at least 2 columns "
                f"(category + at least one programs column)."
            )
        program_cols_by_page[anchor.page] = num_cols - anchor.topic_start_col

    categories = _sort_programs_table_categories(list(programs_by_category.items()))
    categories_per_page = _paginate_categories_by_program_limit(
        categories,
        anchors,
        max_programs_per_table_page=max_programs_per_table_page,
        single_page_soft_overflow=single_page_soft_overflow,
    )

    for anchor, page_categories in zip(anchors, categories_per_page):
        _ensure_table_rows(anchor.table, anchor.anchor_row, len(page_categories))
        category_font_source = None
        topic_font_source = None
        category_paragraph_source = None
        topic_paragraph_source = None
        if anchor.anchor_row < len(anchor.table.rows):
            category_paragraph_source = anchor.table.cell(
                anchor.anchor_row, anchor.category_col
            ).text_frame.paragraphs[0]
            category_runs = list(category_paragraph_source.runs)
            if category_runs:
                category_font_source = category_runs[0]
            topic_paragraph_source = anchor.table.cell(
                anchor.anchor_row, anchor.topic_start_col
            ).text_frame.paragraphs[0]
            topic_runs = list(topic_paragraph_source.runs)
            if topic_runs:
                topic_font_source = topic_runs[0]

        for row_idx in range(anchor.anchor_row, len(anchor.table.rows)):
            for col_idx in range(len(anchor.table.columns)):
                anchor.table.cell(row_idx, col_idx).text = ""

        row_indices = list(
            range(anchor.anchor_row, anchor.anchor_row + len(page_categories))
        )
        programs_cols = program_cols_by_page[anchor.page]

        for row_idx, (category_name, programs) in zip(row_indices, page_categories):
            if not category_name:
                break

            _set_table_cell_lines(
                anchor.table.cell(row_idx, anchor.category_col),
                [category_name],
                bullet=False,
                font_source_run=category_font_source,
                paragraph_source=category_paragraph_source,
            )
            truncated_programs = [
                _truncate_text(program, topic_max_chars, topic_truncation_suffix)
                for program in programs
            ]
            split_programs = _split_evenly(truncated_programs, programs_cols)
            for col_offset, chunk in enumerate(
                split_programs, start=anchor.topic_start_col
            ):
                _set_table_cell_lines(
                    anchor.table.cell(row_idx, col_offset),
                    chunk,
                    bullet=True,
                    font_source_run=topic_font_source,
                    paragraph_source=topic_paragraph_source,
                )

    for anchor, page_categories in zip(anchors, categories_per_page):
        if not page_categories and anchor.page > 1:
            _remove_slide(prs, anchor.slide)

    return True


def _replace_placeholders_in_presentation(
    prs: Presentation, replacements: dict[str, str]
) -> None:
    """Replace all {{placeholder}} patterns in the presentation text.

    Preserves original formatting (font, size, color, bold, italic, etc.)
    by processing runs individually.

    Warns about:
    - Replacement keys that have no matching placeholder in the template
    - Placeholders in the template that have no matching replacement key
    """
    # Scan template text before replacing
    template_placeholders = _extract_template_placeholders(prs)
    replacement_keys = set(replacements.keys())

    unused_keys = replacement_keys - template_placeholders
    if unused_keys:
        log.warning(
            f"Computed {len(unused_keys)} replacement(s) with no matching "
            f"placeholder in the template: {sorted(unused_keys)}"
        )

    missing_keys = template_placeholders - replacement_keys
    if missing_keys:
        log.warning(
            f"Template has {len(missing_keys)} placeholder(s) with no "
            f"replacement value: {sorted(missing_keys)}"
        )

    for slide in prs.slides:
        for shape in _iter_shapes(slide.shapes):
            # Process table cells
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            _replace_placeholders_in_paragraph(paragraph, replacements)
                continue

            if not shape.has_text_frame:
                continue
            text_frame = shape.text_frame

            # Process each paragraph
            for paragraph in text_frame.paragraphs:
                _replace_placeholders_in_paragraph(paragraph, replacements)


def generate_presentation(
    company: str,
    start: str,
    end: str,
    metrics: Metrics,
    output_path: Path,
    max_topics: int,
    max_programs_per_table_page: int,
    single_page_soft_overflow: int,
    topic_max_chars: int | None,
    topic_truncation_suffix: str | None,
    template_path: Path,
) -> Path:
    """Load template, replace placeholders with metrics data, and save to output_path."""
    if not template_path.exists():
        raise FileNotFoundError(f"Template not found: {template_path}")

    log.info(f"Loading template from: {template_path}")
    prs = Presentation(str(template_path))

    replacements = {
        "company": company,
        "period": _format_date_range(start, end),
        "period_phrase": _format_period_phrase(start, end),
        "first_sign_up_date": metrics.first_sign_up_date,
        "last_sign_up_date": metrics.last_sign_up_date,
        "total_sign_ups": str(metrics.total_sign_ups),
        "withdrawn_sign_ups": str(metrics.withdrawn_sign_ups),
        "total_attended_sign_ups": str(metrics.total_attended_sign_ups),
        "total_evaluation_count": str(metrics.total_evaluation_count),
        "cost_per_training": _format_decimal(metrics.cost_per_training),
        "total_training_cost": _format_currency(metrics.total_training_cost),
        "first_attended_date": metrics.first_attended_date,
        "first_attended_topic": metrics.first_attended_topic,
        "last_attended_date": metrics.last_attended_date,
        "last_attended_topic": metrics.last_attended_topic,
        "total_sessions": str(metrics.total_sessions),
        "unique_topics": str(metrics.unique_topics),
        "unique_attendees": str(metrics.unique_attendees),
        "date_slideshow_created": datetime.now().strftime("%B %d, %Y"),
    }

    filled_top_programs_table = _populate_top_programs_table(prs, metrics.top_topics)

    # Expand top topics into top_1_name, top_1_count, etc.
    # Fill missing slots with empty strings so leftover placeholders don't show.
    if not filled_top_programs_table:
        for i in range(1, max_topics + 1):
            if i <= len(metrics.top_topics):
                name, count = metrics.top_topics[i - 1]
                replacements[f"top_{i}_name"] = name
                replacements[f"top_{i}_count"] = (
                    f"{count} sign-up{'s' if count != 1 else ''}"
                )
            else:
                replacements[f"top_{i}_name"] = ""
                replacements[f"top_{i}_count"] = ""

    filled_program_tables = _populate_programs_tables(
        prs,
        metrics.programs_by_category,
        max_programs_per_table_page=max_programs_per_table_page,
        single_page_soft_overflow=single_page_soft_overflow,
        topic_max_chars=topic_max_chars,
        topic_truncation_suffix=topic_truncation_suffix,
    )
    if not filled_program_tables:
        template_placeholders = _extract_template_placeholders(prs)
        replacements.update(
            _build_program_column_replacements(
                metrics.programs_by_category,
                template_placeholders,
                topic_max_chars=topic_max_chars,
                topic_truncation_suffix=topic_truncation_suffix,
            )
        )

    if not filled_top_programs_table:
        _prune_top_table_rows(
            prs,
            top_count=min(len(metrics.top_topics), max_topics),
        )

    _replace_placeholders_in_presentation(prs, replacements)

    # Ensure output directory exists
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # Save presentation
    prs.save(str(output_path))
    log.info(f"Generated presentation: {output_path}")

    return output_path


def _format_period_phrase(start: str, end: str) -> str:
    """Format period with the appropriate preposition for sentence use."""
    start_dt = datetime.strptime(start, "%Y-%m-%d")
    end_dt = datetime.strptime(end, "%Y-%m-%d")

    start_month = start_dt.strftime("%B")
    end_month = end_dt.strftime("%B")
    start_year = start_dt.strftime("%Y")
    end_year = end_dt.strftime("%Y")

    if start_month == end_month and start_year == end_year:
        return f"in {start_month} {start_year}"
    if start_year == end_year:
        return f"between {start_month} and {end_month} {start_year}"
    return f"between {start_month} {start_year} and {end_month} {end_year}"
