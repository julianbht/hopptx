from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_THEME_COLOR
from pptx.enum.text import PP_ALIGN

from hopptx.generator import (
    _format_date_range,
    _format_period_phrase,
    _replace_placeholders_in_presentation,
    _replace_placeholders_in_paragraph,
    _set_table_cell_lines,
    _sort_programs_table_categories,
)


def test_split_placeholder_preserves_theme_font_color() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    textbox = slide.shapes.add_textbox(left=0, top=0, width=1, height=1)
    paragraph = textbox.text_frame.paragraphs[0]

    first_run = paragraph.add_run()
    first_run.text = "{{date_"
    first_run.font.color.theme_color = MSO_THEME_COLOR.BACKGROUND_1

    second_run = paragraph.add_run()
    second_run.text = "slideshow_created}}"

    _replace_placeholders_in_paragraph(
        paragraph, {"date_slideshow_created": "April 25, 2026"}
    )

    assert len(paragraph.runs) == 1
    replaced_run = paragraph.runs[0]
    assert replaced_run.text == "April 25, 2026"
    assert replaced_run.font.color.type == MSO_COLOR_TYPE.SCHEME
    assert replaced_run.font.color.theme_color == MSO_THEME_COLOR.BACKGROUND_1


def test_set_table_cell_lines_preserves_paragraph_alignment() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table_shape = slide.shapes.add_table(rows=1, cols=1, left=0, top=0, width=1, height=1)
    cell = table_shape.table.cell(0, 0)

    source_paragraph = cell.text_frame.paragraphs[0]
    source_paragraph.alignment = PP_ALIGN.CENTER
    source_run = source_paragraph.add_run()
    source_run.text = "template"

    _set_table_cell_lines(
        cell,
        ["42"],
        bullet=False,
        font_source_run=source_run,
        paragraph_source=source_paragraph,
    )

    assert cell.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER


def test_format_date_range_collapses_same_month() -> None:
    assert _format_date_range("2026-04-01", "2026-04-30") == "April 2026"


def test_format_period_phrase_for_single_month() -> None:
    assert _format_period_phrase("2026-04-01", "2026-04-30") == "in April 2026"


def test_replace_placeholders_in_group_shape() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    textbox = group.shapes.add_textbox(left=0, top=0, width=1, height=1)
    textbox.text_frame.paragraphs[0].text = "{{period_phrase}}"

    _replace_placeholders_in_presentation(prs, {"period_phrase": "in April 2026"})

    assert textbox.text_frame.paragraphs[0].text == "in April 2026"


def test_sort_programs_table_categories_uses_fixed_order() -> None:
    categories = [
        ("Influencing Skills", ["I1"]),
        ("Business Skills", ["B1"]),
        ("Personal / Professional Development", ["P1"]),
        ("Technical Skills", ["T1"]),
        ("Communication Skills", ["C1"]),
        ("Leading Others", ["L1"]),
    ]

    sorted_categories = _sort_programs_table_categories(categories)

    assert [name for name, _ in sorted_categories] == [
        "Technical Skills",
        "Business Skills",
        "Leading Others",
        "Communication Skills",
        "Influencing Skills",
        "Personal / Professional Development",
    ]


def test_sort_programs_table_categories_keeps_unknowns_after_known() -> None:
    categories = [
        ("Custom Category", ["X1"]),
        ("Communication Skills", ["C1"]),
        ("Technical Skills", ["T1"]),
        ("Another Custom", ["Y1"]),
    ]

    sorted_categories = _sort_programs_table_categories(categories)

    assert [name for name, _ in sorted_categories] == [
        "Technical Skills",
        "Communication Skills",
        "Custom Category",
        "Another Custom",
    ]
